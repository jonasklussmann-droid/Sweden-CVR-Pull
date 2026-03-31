"""Export financial data to JSON and Excel formats."""

import json
import logging
from pathlib import Path
from typing import Optional

from .models import FinancialData

logger = logging.getLogger(__name__)

# Column headers for Excel export (human-readable)
EXCEL_COLUMNS = [
    ("company_name", "Company Name"),
    ("org_number", "Org Number"),
    ("nace_code", "NACE/SNI Code"),
    ("fiscal_year", "Fiscal Year"),
    ("period_start", "Period Start"),
    ("period_end", "Period End"),
    ("currency", "Currency"),
    ("revenue", "Revenue"),
    ("net_sales", "Net Sales"),
    ("other_operating_income", "Other Operating Income"),
    ("change_in_inventory", "Change in Inventory"),
    ("raw_materials_and_consumables", "Raw Materials & Consumables"),
    ("trade_goods_cogs", "Trade Goods (COGS)"),
    ("other_external_costs", "Other External Costs"),
    ("personnel_costs", "Personnel Costs"),
    ("depreciation_amortization", "Depreciation & Amortization"),
    ("other_operating_expenses", "Other Operating Expenses"),
    ("operating_profit", "Operating Profit"),
    ("financial_income", "Financial Income"),
    ("financial_costs", "Financial Costs"),
    ("result_after_financial_items", "Result After Financial Items"),
    ("tax", "Tax"),
    ("net_income", "Net Income"),
    ("cogs", "COGS (Calculated)"),
    ("gross_profit", "Gross Profit (Calculated)"),
    ("ebit", "EBIT"),
    ("ebitda", "EBITDA (Calculated)"),
    ("source_filing", "Source Filing"),
    ("source_file", "Source File"),
]


def export_to_json(
    data: list[FinancialData],
    output_path: str = "output.json",
    indent: int = 2,
) -> str:
    """Export financial data to a JSON file.

    Args:
        data: List of FinancialData records
        output_path: Output file path
        indent: JSON indentation level

    Returns:
        Path to the created JSON file
    """
    records = []
    for financial in data:
        record = financial.to_dict()
        # Add individual fact traceability
        record["_facts"] = [
            {
                "concept": f.concept,
                "value": f.value,
                "text_value": f.text_value,
                "context_id": f.context_id,
                "unit": f.unit,
                "period_start": f.period_start,
                "period_end": f.period_end,
                "scale": f.scale,
            }
            for f in financial.facts
            if f.value is not None  # Only include numeric facts for traceability
        ]
        records.append(record)

    output = {
        "metadata": {
            "total_records": len(records),
            "export_format": "json",
            "schema_version": "1.0",
        },
        "data": records,
    }

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=indent, ensure_ascii=False, default=str)

    logger.info("Exported %d records to %s", len(records), output_path)
    return str(path)


def export_to_excel(
    data: list[FinancialData],
    output_path: str = "output.xlsx",
) -> str:
    """Export financial data to an Excel file.

    Args:
        data: List of FinancialData records
        output_path: Output file path

    Returns:
        Path to the created Excel file
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, numbers
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Financial Data"

    # Style definitions
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=11, color="FFFFFF")
    number_format = '#,##0'
    decimal_format = '#,##0.00'

    # Write headers
    for col_idx, (field_key, header_label) in enumerate(EXCEL_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_label)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    # Write data rows
    for row_idx, financial in enumerate(data, 2):
        record = financial.to_dict()
        for col_idx, (field_key, _) in enumerate(EXCEL_COLUMNS, 1):
            value = record.get(field_key)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)

            # Apply number formatting for financial columns
            if isinstance(value, (int, float)) and value is not None:
                cell.number_format = number_format
                cell.alignment = Alignment(horizontal="right")

    # Auto-adjust column widths
    for col_idx in range(1, len(EXCEL_COLUMNS) + 1):
        col_letter = get_column_letter(col_idx)
        max_length = len(EXCEL_COLUMNS[col_idx - 1][1])  # Start with header length
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 3, 30)

    # Freeze the header row
    ws.freeze_panes = "A2"

    # Add auto-filter
    ws.auto_filter.ref = ws.dimensions

    # Save
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))

    logger.info("Exported %d records to %s", len(data), output_path)
    return str(path)


def export_all(
    data: list[FinancialData],
    output_dir: str = "output",
    base_name: str = "swedish_financials",
) -> dict[str, str]:
    """Export financial data to both JSON and Excel formats.

    Args:
        data: List of FinancialData records
        output_dir: Output directory
        base_name: Base filename (without extension)

    Returns:
        Dictionary with format -> file path mappings
    """
    Path(output_dir).mkdir(parents=True, exist_ok=True)

    json_path = export_to_json(data, f"{output_dir}/{base_name}.json")
    excel_path = export_to_excel(data, f"{output_dir}/{base_name}.xlsx")

    return {"json": json_path, "excel": excel_path}
